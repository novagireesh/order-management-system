"""FastAPI server for the Order Management System. Run with start.bat or: python server.py"""
import io
import os
import re
import sys
import threading
import webbrowser
from datetime import datetime
from typing import Optional

import openpyxl
from fastapi import FastAPI, File, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import auth  # noqa: E402
import pdf  # noqa: E402
from db import (CONTRACTOR_TYPES, IMAGE_DIR, PDF_DIR, STAGES, fetch_order_items, get_db, get_settings,  # noqa: E402
                init_db, next_pi_number, next_po_number, row, rows, set_setting, today)

STATIC_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "static")
app = FastAPI(title="Order Management System")
app.add_middleware(auth.AuthMiddleware)
app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")
app.mount("/images", StaticFiles(directory=IMAGE_DIR), name="images")
app.mount("/pdf", StaticFiles(directory=PDF_DIR), name="pdf")

init_db()
auth.init_auth()


def err(msg, code=400):
    raise HTTPException(status_code=code, detail=msg)


def f(v, default=0.0):
    try:
        return float(v) if v not in (None, "") else default
    except (TypeError, ValueError):
        return default


def clean(s):
    return re.sub(r"\s+", " ", str(s or "")).strip()


@app.get("/", response_class=HTMLResponse)
def index():
    with open(os.path.join(STATIC_DIR, "index.html"), encoding="utf8") as fh:
        return fh.read()


# ----------------------------------------------------------------------------- auth / users

_login_fails = {}  # ip -> [timestamps]; simple brute-force protection


@app.get("/api/brand")
def api_brand():
    with get_db() as conn:
        s = get_settings(conn)
        return {"company_name": s.get("company_name", ""), "company_short": s.get("company_short") or s.get("company_name", "")}


@app.post("/api/login")
async def api_login(req: Request):
    import time
    data = await req.json()
    ip = req.headers.get("x-forwarded-for", req.client.host if req.client else "?").split(",")[0].strip()
    fails = [t for t in _login_fails.get(ip, []) if time.time() - t < 900]
    if len(fails) >= 10:
        err("Too many failed attempts. Try again in 15 minutes.", 429)
    user, token = auth.login(data.get("username", ""), data.get("password", ""), req.headers.get("user-agent", ""))
    if not user:
        _login_fails[ip] = fails + [time.time()]
        err("Wrong username or password", 401)
    _login_fails.pop(ip, None)
    resp = JSONResponse(user)
    secure = req.headers.get("x-forwarded-proto", req.url.scheme) == "https"
    resp.set_cookie(auth.COOKIE, token, max_age=auth.SESSION_DAYS * 86400, httponly=True, samesite="lax", secure=secure)
    return resp


@app.post("/api/logout")
def api_logout(req: Request):
    auth.logout(req.cookies.get(auth.COOKIE))
    resp = JSONResponse({"ok": True})
    resp.delete_cookie(auth.COOKIE)
    return resp


@app.get("/api/me")
def api_me(req: Request):
    u = auth.user_from_token(req.cookies.get(auth.COOKIE))
    if not u:
        err("Not logged in", 401)
    return u


@app.put("/api/me/password")
async def api_me_password(req: Request):
    data = await req.json()
    return auth.change_password(auth.require_user(req)["id"], data.get("old"), data.get("new"))


@app.get("/api/users")
def api_users():
    return {"users": auth.list_users(), "roles": [{"id": r, "label": auth.ROLE_LABELS[r]} for r in auth.ROLES]}


@app.post("/api/users")
async def api_user_create(req: Request):
    d = await req.json()
    return auth.create_user(d.get("username"), d.get("name"), d.get("role"), d.get("password"))


@app.put("/api/users/{id}")
async def api_user_update(id: int, req: Request):
    return auth.update_user(id, await req.json(), auth.require_user(req))


# ----------------------------------------------------------------------------- meta / settings

@app.get("/api/meta")
def meta(request: Request):
    with get_db() as conn:
        s = get_settings(conn)
        return {
            "user": auth.user_from_token(request.cookies.get(auth.COOKIE)),
            "stages": STAGES,
            "contractor_types": CONTRACTOR_TYPES,
            "settings": s,
            "buyers": rows(conn.execute("SELECT * FROM buyers WHERE active=1 ORDER BY name")),
            "contractors": rows(conn.execute("SELECT * FROM contractors WHERE active=1 ORDER BY type, name")),
            "next_po": _peek_po(s),
        }


def _peek_po(s):
    return f"{s.get('po_prefix','CPO-')}{str(int(s.get('po_next_number','1') or 1)).zfill(int(s.get('po_number_width','3') or 3))}{s.get('po_suffix','')}"


@app.put("/api/settings")
async def update_settings(req: Request):
    data = await req.json()
    with get_db() as conn:
        for k, v in data.items():
            set_setting(conn, k, v)
        return get_settings(conn)


# ----------------------------------------------------------------------------- dashboard

@app.get("/api/dashboard")
def dashboard():
    with get_db() as conn:
        items = fetch_order_items(conn, "oi.qty > 0")
        pending = [i for i in items if i["remaining_qty"] > 0]
        by_buyer = {}
        for i in pending:
            b = by_buyer.setdefault(i["buyer_name"], {"buyer": i["buyer_name"], "qty": 0, "cbm": 0, "amount": 0, "pos": set()})
            b["qty"] += i["remaining_qty"]
            b["cbm"] += i["remaining_cbm"]
            b["amount"] += i["remaining_amount"]
            b["pos"].add(i["po_number"])
        for b in by_buyer.values():
            b["pos"] = len(b["pos"])
            b["cbm"] = round(b["cbm"], 3)
            b["amount"] = round(b["amount"], 2)
        stage_counts = rows(conn.execute(
            """SELECT pi.stage, COUNT(*) AS lines, SUM(pi.qty) AS qty FROM contractor_po_items pi
               JOIN order_items oi ON oi.id = pi.order_item_id
               WHERE (oi.qty - COALESCE((SELECT SUM(qty) FROM dispatch_items d WHERE d.order_item_id=oi.id),0)) > 0
               GROUP BY pi.stage ORDER BY pi.stage"""))
        contractor_load = rows(conn.execute(
            """SELECT c.name, c.type, COUNT(pi.id) AS lines, SUM(pi.qty) AS qty, SUM(pi.qty*pi.labour_rate) AS labour,
                      SUM(pi.qty*oi.cbm) AS cbm
               FROM contractor_po_items pi JOIN contractor_pos cp ON cp.id=pi.po_id JOIN contractors c ON c.id=cp.contractor_id
               JOIN order_items oi ON oi.id=pi.order_item_id
               WHERE pi.stage < 11 AND (oi.qty - COALESCE((SELECT SUM(qty) FROM dispatch_items d WHERE d.order_item_id=oi.id),0)) > 0
               GROUP BY c.id ORDER BY qty DESC"""))
        unassigned = [i for i in pending if i["remaining_to_assign"] > 0]
        recent_docs = rows(conn.execute("SELECT * FROM documents ORDER BY id DESC LIMIT 10"))
        return {
            "open_orders": conn.execute("SELECT COUNT(*) FROM orders WHERE status='open'").fetchone()[0],
            "pending_items": len(pending),
            "pending_qty": sum(i["remaining_qty"] for i in pending),
            "pending_cbm": round(sum(i["remaining_cbm"] for i in pending), 3),
            "pending_amount": round(sum(i["remaining_amount"] for i in pending), 2),
            "unassigned_lines": len(unassigned),
            "unassigned_qty": sum(i["remaining_to_assign"] for i in unassigned),
            "by_buyer": sorted(by_buyer.values(), key=lambda b: -b["cbm"]),
            "stages": [{"stage": s["stage"], "name": STAGES[s["stage"]] if s["stage"] is not None else "?", "lines": s["lines"], "qty": s["qty"]} for s in stage_counts],
            "contractor_load": contractor_load,
            "recent_docs": recent_docs,
        }


# ----------------------------------------------------------------------------- buyers / contractors

def _crud(table, fields):
    def list_():
        with get_db() as conn:
            return rows(conn.execute(f"SELECT * FROM {table} ORDER BY name"))

    async def create(req: Request):
        data = await req.json()
        if not clean(data.get("name")):
            err("Name is required")
        with get_db() as conn:
            cols = ["name"] + [k for k in fields if k in data]
            vals = [clean(data["name"])] + [data.get(k) for k in fields if k in data]
            try:
                conn.execute(f"INSERT INTO {table}({','.join(cols)}) VALUES ({','.join('?'*len(cols))})", vals)
            except Exception as e:
                err(f"Could not save: {e}")
            return row(conn.execute(f"SELECT * FROM {table} WHERE id=last_insert_rowid()"))

    async def update(id: int, req: Request):
        data = await req.json()
        with get_db() as conn:
            cols = [k for k in ["name"] + fields + ["active"] if k in data]
            if not cols:
                return row(conn.execute(f"SELECT * FROM {table} WHERE id=?", (id,)))
            try:
                conn.execute(f"UPDATE {table} SET {','.join(k+'=?' for k in cols)} WHERE id=?", [clean(data[k]) if k == 'name' else data[k] for k in cols] + [id])
            except Exception as e:
                err(f"Could not save: {e}")
            return row(conn.execute(f"SELECT * FROM {table} WHERE id=?", (id,)))

    def delete(id: int):
        with get_db() as conn:
            try:
                conn.execute(f"DELETE FROM {table} WHERE id=?", (id,))
            except Exception:
                conn.execute(f"UPDATE {table} SET active=0 WHERE id=?", (id,))
                return {"ok": True, "note": "In use - marked inactive instead of deleting"}
            return {"ok": True}

    return list_, create, update, delete


b_list, b_create, b_update, b_delete = _crud("buyers", ["address", "email", "phone", "country", "notes"])
app.get("/api/buyers")(b_list)
app.post("/api/buyers")(b_create)
app.put("/api/buyers/{id}")(b_update)
app.delete("/api/buyers/{id}")(b_delete)
c_list, c_create, c_update, c_delete = _crud("contractors", ["type", "address", "phone", "notes"])
app.get("/api/contractors")(c_list)
app.post("/api/contractors")(c_create)
app.put("/api/contractors/{id}")(c_update)
app.delete("/api/contractors/{id}")(c_delete)


# ----------------------------------------------------------------------------- products

@app.get("/api/products")
def products(q: str = "", buyer_id: Optional[int] = None, limit: int = 2000):
    with get_db() as conn:
        where, params = ["1=1"], []
        if q:
            where.append("(p.code LIKE ? OR p.name LIKE ? OR p.internal_code LIKE ?)")
            params += [f"%{q}%"] * 3
        if buyer_id:
            where.append("p.buyer_id=?")
            params.append(buyer_id)
        sql = f"""SELECT p.*, b.name AS buyer_name,
                  (SELECT COUNT(*) FROM order_items oi WHERE oi.product_id=p.id) AS times_ordered
                  FROM products p LEFT JOIN buyers b ON b.id=p.buyer_id WHERE {' AND '.join(where)} ORDER BY p.code LIMIT ?"""
        return rows(conn.execute(sql, params + [limit]))


@app.get("/api/products/lookup")
def product_lookup(code: str):
    with get_db() as conn:
        p = row(conn.execute("SELECT p.*, b.name AS buyer_name FROM products p LEFT JOIN buyers b ON b.id=p.buyer_id WHERE UPPER(p.code)=UPPER(?)", (clean(code),)))
        if not p:
            return JSONResponse({"found": False})
        p["found"] = True
        return p


PRODUCT_FIELDS = ["code", "name", "internal_code", "size", "cbm", "rate", "labour_rate", "finish", "buyer_id", "notes", "active"]


@app.post("/api/products")
async def product_create(req: Request):
    data = await req.json()
    code = clean(data.get("code")).upper()
    if not code:
        err("Product code is required")
    with get_db() as conn:
        if conn.execute("SELECT 1 FROM products WHERE code=?", (code,)).fetchone():
            err(f"Product {code} already exists")
        conn.execute("INSERT INTO products(code,name,internal_code,size,cbm,rate,labour_rate,finish,buyer_id,notes) VALUES (?,?,?,?,?,?,?,?,?,?)",
                     (code, clean(data.get("name")), clean(data.get("internal_code")), clean(data.get("size")), f(data.get("cbm")), f(data.get("rate")),
                      f(data.get("labour_rate")), clean(data.get("finish")), data.get("buyer_id") or None, data.get("notes") or ""))
        return row(conn.execute("SELECT * FROM products WHERE code=?", (code,)))


@app.put("/api/products/{id}")
async def product_update(id: int, req: Request):
    data = await req.json()
    with get_db() as conn:
        cols, vals = [], []
        for k in PRODUCT_FIELDS:
            if k in data:
                v = data[k]
                if k in ("cbm", "rate", "labour_rate"):
                    v = f(v)
                elif k == "code":
                    v = clean(v).upper()
                elif k == "buyer_id":
                    v = v or None
                elif k != "active":
                    v = clean(v) if k != "notes" else (v or "")
                cols.append(f"{k}=?")
                vals.append(v)
        if cols:
            try:
                conn.execute(f"UPDATE products SET {','.join(cols)} WHERE id=?", vals + [id])
            except Exception as e:
                err(f"Could not save: {e}")
        return row(conn.execute("SELECT * FROM products WHERE id=?", (id,)))


@app.delete("/api/products/{id}")
def product_delete(id: int):
    with get_db() as conn:
        used = conn.execute("SELECT COUNT(*) FROM order_items WHERE product_id=?", (id,)).fetchone()[0]
        if used:
            conn.execute("UPDATE products SET active=0 WHERE id=?", (id,))
            return {"ok": True, "note": "Product is used in orders - marked inactive"}
        conn.execute("DELETE FROM products WHERE id=?", (id,))
        return {"ok": True}


@app.post("/api/products/{id}/image")
async def product_image(id: int, file: UploadFile = File(...)):
    from PIL import Image
    with get_db() as conn:
        p = row(conn.execute("SELECT * FROM products WHERE id=?", (id,)))
        if not p:
            err("Product not found", 404)
        fname = re.sub(r"[^A-Za-z0-9_-]+", "_", p["code"]) + ".jpg"
        data = await file.read()
        im = Image.open(io.BytesIO(data)).convert("RGB")
        im.thumbnail((600, 600))
        im.save(os.path.join(IMAGE_DIR, fname), "JPEG", quality=85)
        conn.execute("UPDATE products SET image=? WHERE id=?", (fname, id))
        return {"image": fname}


# ----------------------------------------------------------------------------- orders

def _order_full(conn, oid):
    o = row(conn.execute("SELECT o.*, b.name AS buyer_name FROM orders o JOIN buyers b ON b.id=o.buyer_id WHERE o.id=?", (oid,)))
    if not o:
        err("Order not found", 404)
    o["items"] = fetch_order_items(conn, "oi.order_id=?", (oid,))
    o["total_qty"] = sum(i["qty"] for i in o["items"])
    o["total_cbm"] = round(sum(i["t_cbm"] for i in o["items"]), 3)
    o["total_amount"] = round(sum(i["amount"] for i in o["items"]), 2)
    o["total_labour"] = round(sum(i["labour_amount"] for i in o["items"]), 2)
    o["remaining_qty"] = sum(i["remaining_qty"] for i in o["items"])
    o["remaining_cbm"] = round(sum(i["remaining_cbm"] for i in o["items"]), 3)
    o["remaining_amount"] = round(sum(i["remaining_amount"] for i in o["items"]), 2)
    o["documents"] = rows(conn.execute("SELECT * FROM documents WHERE ref_type='order' AND ref_id=? ORDER BY id DESC", (oid,)))
    o["dispatches"] = rows(conn.execute(
        "SELECT d.*, (SELECT SUM(qty) FROM dispatch_items di WHERE di.dispatch_id=d.id) AS qty FROM dispatches d WHERE d.order_id=? ORDER BY d.dispatch_date, d.id", (oid,)))
    return o


def _generate_order_docs(conn, oid, kinds=("pi", "carton", "sticker")):
    o = _order_full(conn, oid)
    settings = get_settings(conn)
    buyer = row(conn.execute("SELECT * FROM buyers WHERE id=?", (o["buyer_id"],)))
    o["doc_date"] = today()
    out = []
    for kind in kinds:
        if kind == "pi":
            if not o.get("pi_number"):
                o["pi_number"] = next_pi_number(conn)
                conn.execute("UPDATE orders SET pi_number=? WHERE id=?", (o["pi_number"], oid))
            fname = pdf.pi_pdf(o, o["items"], buyer, settings)
        else:
            fname = pdf.carton_sticker_pdf(kind, o, o["items"], settings)
            col = "carton_ordered_at" if kind == "carton" else "sticker_ordered_at"
            conn.execute(f"UPDATE order_items SET {col}=COALESCE({col}, ?) WHERE order_id=?", (today(), oid))
        conn.execute("DELETE FROM documents WHERE ref_type='order' AND ref_id=? AND kind=?", (oid, kind))
        conn.execute("INSERT INTO documents(kind, ref_type, ref_id, filename) VALUES (?,?,?,?)", (kind, "order", oid, fname))
        out.append({"kind": kind, "filename": fname})
    return out


@app.get("/api/orders")
def orders(buyer_id: Optional[int] = None, po: str = "", status: str = "", date_from: str = "", date_to: str = "", q: str = ""):
    with get_db() as conn:
        where, params = ["1=1"], []
        if buyer_id:
            where.append("o.buyer_id=?"); params.append(buyer_id)
        if po:
            where.append("o.po_number LIKE ?"); params.append(f"%{po}%")
        if status:
            where.append("o.status=?"); params.append(status)
        if date_from:
            where.append("o.order_date >= ?"); params.append(date_from)
        if date_to:
            where.append("o.order_date <= ?"); params.append(date_to)
        if q:
            where.append("(o.po_number LIKE ? OR b.name LIKE ? OR EXISTS(SELECT 1 FROM order_items x WHERE x.order_id=o.id AND x.code LIKE ?))")
            params += [f"%{q}%"] * 3
        sql = f"""SELECT o.*, b.name AS buyer_name,
                 (SELECT COUNT(*) FROM order_items oi WHERE oi.order_id=o.id) AS lines,
                 (SELECT SUM(qty) FROM order_items oi WHERE oi.order_id=o.id) AS total_qty,
                 (SELECT SUM(qty*cbm) FROM order_items oi WHERE oi.order_id=o.id) AS total_cbm,
                 (SELECT SUM(qty*rate) FROM order_items oi WHERE oi.order_id=o.id) AS total_amount,
                 (SELECT SUM(qty*labour_rate) FROM order_items oi WHERE oi.order_id=o.id) AS total_labour,
                 (SELECT SUM(di.qty) FROM dispatch_items di JOIN order_items oi ON oi.id=di.order_item_id WHERE oi.order_id=o.id) AS dispatched_qty
                 FROM orders o JOIN buyers b ON b.id=o.buyer_id WHERE {' AND '.join(where)}
                 ORDER BY o.order_date DESC, o.id DESC"""
        out = rows(conn.execute(sql, params))
        for o in out:
            o["remaining_qty"] = (o["total_qty"] or 0) - (o["dispatched_qty"] or 0)
            o["total_cbm"] = round(o["total_cbm"] or 0, 3)
            o["total_amount"] = round(o["total_amount"] or 0, 2)
            o["total_labour"] = round(o["total_labour"] or 0, 2)
        return out


@app.get("/api/orders/{id}")
def order_get(id: int):
    with get_db() as conn:
        return _order_full(conn, id)


def _save_items(conn, oid, items):
    conn.execute("DELETE FROM order_items WHERE order_id=? AND id NOT IN (SELECT order_item_id FROM contractor_po_items UNION SELECT order_item_id FROM dispatch_items)", (oid,))
    existing = {r["id"]: r for r in rows(conn.execute("SELECT * FROM order_items WHERE order_id=?", (oid,)))}
    seen = set()
    for n, it in enumerate(items, 1):
        code = clean(it.get("code")).upper()
        qty = f(it.get("qty"))
        if not code or qty <= 0:
            continue
        p = row(conn.execute("SELECT * FROM products WHERE code=?", (code,)))
        if not p:
            # unknown code: create product master so it autofills next time
            conn.execute("INSERT INTO products(code,name,size,cbm,rate,labour_rate,finish,buyer_id) VALUES (?,?,?,?,?,?,?,?)",
                         (code, clean(it.get("name")), clean(it.get("size")), f(it.get("cbm")), f(it.get("rate")), f(it.get("labour_rate")), clean(it.get("finish")),
                          conn.execute("SELECT buyer_id FROM orders WHERE id=?", (oid,)).fetchone()[0]))
            p = row(conn.execute("SELECT * FROM products WHERE code=?", (code,)))
        else:
            # remember labour rate / cbm / rate on master if it was blank
            upd = {}
            for k in ("labour_rate", "cbm", "rate"):
                if not p.get(k) and f(it.get(k)):
                    upd[k] = f(it.get(k))
            if not p.get("size") and clean(it.get("size")):
                upd["size"] = clean(it.get("size"))
            if not p.get("name") and clean(it.get("name")):
                upd["name"] = clean(it.get("name"))
            if upd:
                conn.execute(f"UPDATE products SET {','.join(k+'=?' for k in upd)} WHERE id=?", list(upd.values()) + [p["id"]])
        vals = (p["id"], n, code, clean(it.get("name")) or p["name"], clean(it.get("size")) or p["size"], qty,
                f(it.get("rate"), p["rate"]), f(it.get("cbm"), p["cbm"]), f(it.get("labour_rate"), p["labour_rate"]),
                clean(it.get("finish")) or p["finish"], 1 if it.get("iron_required") else 0, it.get("notes") or "")
        iid = it.get("id")
        if iid and int(iid) in existing:
            conn.execute("""UPDATE order_items SET product_id=?, sno=?, code=?, name=?, size=?, qty=?, rate=?, cbm=?, labour_rate=?, finish=?, iron_required=?, notes=? WHERE id=?""",
                         vals + (int(iid),))
            seen.add(int(iid))
        else:
            conn.execute("""INSERT INTO order_items(order_id, product_id, sno, code, name, size, qty, rate, cbm, labour_rate, finish, iron_required, notes)
                            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)""", (oid,) + vals)
    for iid, r in existing.items():
        if iid not in seen:
            used = conn.execute("SELECT 1 FROM contractor_po_items WHERE order_item_id=? UNION SELECT 1 FROM dispatch_items WHERE order_item_id=?", (iid, iid)).fetchone()
            if used:
                err(f"Item {r['code']} cannot be removed: it is already assigned to a contractor or dispatched")
            conn.execute("DELETE FROM order_items WHERE id=?", (iid,))


@app.post("/api/orders")
async def order_create(req: Request):
    data = await req.json()
    buyer_id = data.get("buyer_id")
    po = clean(data.get("po_number"))
    if not buyer_id or not po:
        err("Buyer and PO number are required")
    if not data.get("items"):
        err("Add at least one item")
    with get_db() as conn:
        if conn.execute("SELECT 1 FROM orders WHERE buyer_id=? AND po_number=?", (buyer_id, po)).fetchone():
            err(f"PO {po} already exists for this buyer")
        conn.execute("INSERT INTO orders(buyer_id, po_number, order_date, destination, notes, pi_number) VALUES (?,?,?,?,?,?)",
                     (buyer_id, po, data.get("order_date") or today(), clean(data.get("destination")), data.get("notes") or "", clean(data.get("pi_number"))))
        oid = conn.execute("SELECT last_insert_rowid()").fetchone()[0]
        _save_items(conn, oid, data["items"])
        docs = _generate_order_docs(conn, oid) if data.get("generate_docs", True) else []
        o = _order_full(conn, oid)
        o["generated"] = docs
        return o


@app.put("/api/orders/{id}")
async def order_update(id: int, req: Request):
    data = await req.json()
    with get_db() as conn:
        o = row(conn.execute("SELECT * FROM orders WHERE id=?", (id,)))
        if not o:
            err("Order not found", 404)
        po = clean(data.get("po_number", o["po_number"]))
        buyer_id = data.get("buyer_id", o["buyer_id"])
        dup = conn.execute("SELECT 1 FROM orders WHERE buyer_id=? AND po_number=? AND id<>?", (buyer_id, po, id)).fetchone()
        if dup:
            err(f"PO {po} already exists for this buyer")
        conn.execute("UPDATE orders SET buyer_id=?, po_number=?, order_date=?, destination=?, notes=?, pi_number=?, status=? WHERE id=?",
                     (buyer_id, po, data.get("order_date", o["order_date"]), clean(data.get("destination", o["destination"])),
                      data.get("notes", o["notes"]), clean(data.get("pi_number", o["pi_number"])), data.get("status", o["status"]), id))
        if "items" in data:
            _save_items(conn, id, data["items"])
        docs = _generate_order_docs(conn, id) if data.get("generate_docs") else []
        o = _order_full(conn, id)
        o["generated"] = docs
        return o


@app.post("/api/orders/{id}/documents")
def order_docs(id: int, kinds: str = "pi,carton,sticker"):
    with get_db() as conn:
        return _generate_order_docs(conn, id, tuple(k for k in kinds.split(",") if k))


@app.delete("/api/orders/{id}")
def order_delete(id: int):
    with get_db() as conn:
        used = conn.execute("SELECT COUNT(*) FROM contractor_po_items pi JOIN order_items oi ON oi.id=pi.order_item_id WHERE oi.order_id=?", (id,)).fetchone()[0]
        if used:
            err("Order has contractor assignments. Delete those POs first.")
        conn.execute("DELETE FROM documents WHERE ref_type='order' AND ref_id=?", (id,))
        conn.execute("DELETE FROM orders WHERE id=?", (id,))
        return {"ok": True}


@app.get("/api/order-items")
def order_items(buyer_id: Optional[int] = None, order_id: Optional[int] = None, pending: int = 0, unassigned: int = 0, q: str = ""):
    """Flat list of order lines with computed remaining/assigned quantities."""
    with get_db() as conn:
        where, params = ["1=1"], []
        if buyer_id:
            where.append("o.buyer_id=?"); params.append(buyer_id)
        if order_id:
            where.append("o.id=?"); params.append(order_id)
        if q:
            where.append("(oi.code LIKE ? OR oi.name LIKE ? OR o.po_number LIKE ?)"); params += [f"%{q}%"] * 3
        items = fetch_order_items(conn, " AND ".join(where), params)
        if pending:
            items = [i for i in items if i["remaining_qty"] > 0]
        if unassigned:
            items = [i for i in items if i["remaining_to_assign"] > 0 and i["remaining_qty"] > 0]
        return items


# ----------------------------------------------------------------------------- contractor POs / assignment

def _po_full(conn, pid):
    po = row(conn.execute("SELECT cp.*, c.name AS contractor_name, c.type AS contractor_type FROM contractor_pos cp JOIN contractors c ON c.id=cp.contractor_id WHERE cp.id=?", (pid,)))
    if not po:
        err("PO not found", 404)
    po["items"] = rows(conn.execute(
        """SELECT pi.*, oi.code, oi.name, oi.size, oi.cbm, oi.finish, oi.qty AS order_qty, o.po_number AS buyer_po, o.po_number, b.name AS buyer_name, p.image,
                  (SELECT GROUP_CONCAT(stage || '|' || COALESCE(log_date,'') || '|' || COALESCE(note,''), ';;') FROM stage_log sl WHERE sl.po_item_id=pi.id) AS history
           FROM contractor_po_items pi JOIN order_items oi ON oi.id=pi.order_item_id JOIN orders o ON o.id=oi.order_id
           JOIN buyers b ON b.id=o.buyer_id LEFT JOIN products p ON p.id=oi.product_id WHERE pi.po_id=? ORDER BY pi.id""", (pid,)))
    for it in po["items"]:
        it["t_cbm"] = round(it["qty"] * (it["cbm"] or 0), 3)
        it["amount"] = round(it["qty"] * (it["labour_rate"] or 0), 2)
        it["stage_name"] = STAGES[it["stage"]] if it["stage"] is not None and 0 <= it["stage"] < len(STAGES) else ""
    po["total_qty"] = sum(i["qty"] for i in po["items"])
    po["total_cbm"] = round(sum(i["t_cbm"] for i in po["items"]), 3)
    po["total_amount"] = round(sum(i["amount"] for i in po["items"]), 2)
    po["documents"] = rows(conn.execute("SELECT * FROM documents WHERE ref_type='po' AND ref_id=? ORDER BY id DESC", (pid,)))
    return po


def _generate_po_doc(conn, pid):
    po = _po_full(conn, pid)
    contractor = row(conn.execute("SELECT * FROM contractors WHERE id=?", (po["contractor_id"],)))
    fname = pdf.contractor_po_pdf(po, contractor, po["items"], get_settings(conn))
    conn.execute("DELETE FROM documents WHERE ref_type='po' AND ref_id=?", (pid,))
    conn.execute("INSERT INTO documents(kind, ref_type, ref_id, filename) VALUES (?,?,?,?)", ("po", "po", pid, fname))
    return fname


@app.get("/api/pos")
def pos(contractor_id: Optional[int] = None, buyer_id: Optional[int] = None, buyer_po: str = "", po: str = "", date_from: str = "", date_to: str = "", ctype: str = ""):
    with get_db() as conn:
        where, params = ["1=1"], []
        if contractor_id:
            where.append("cp.contractor_id=?"); params.append(contractor_id)
        if ctype:
            where.append("c.type=?"); params.append(ctype)
        if po:
            where.append("cp.po_number LIKE ?"); params.append(f"%{po}%")
        if date_from:
            where.append("cp.po_date >= ?"); params.append(date_from)
        if date_to:
            where.append("cp.po_date <= ?"); params.append(date_to)
        if buyer_id:
            where.append("EXISTS(SELECT 1 FROM contractor_po_items pi JOIN order_items oi ON oi.id=pi.order_item_id JOIN orders o ON o.id=oi.order_id WHERE pi.po_id=cp.id AND o.buyer_id=?)")
            params.append(buyer_id)
        if buyer_po:
            where.append("EXISTS(SELECT 1 FROM contractor_po_items pi JOIN order_items oi ON oi.id=pi.order_item_id JOIN orders o ON o.id=oi.order_id WHERE pi.po_id=cp.id AND o.po_number LIKE ?)")
            params.append(f"%{buyer_po}%")
        sql = f"""SELECT cp.*, c.name AS contractor_name, c.type AS contractor_type,
                  (SELECT COUNT(*) FROM contractor_po_items pi WHERE pi.po_id=cp.id) AS lines,
                  (SELECT SUM(pi.qty) FROM contractor_po_items pi WHERE pi.po_id=cp.id) AS total_qty,
                  (SELECT SUM(pi.qty*oi.cbm) FROM contractor_po_items pi JOIN order_items oi ON oi.id=pi.order_item_id WHERE pi.po_id=cp.id) AS total_cbm,
                  (SELECT SUM(pi.qty*pi.labour_rate) FROM contractor_po_items pi WHERE pi.po_id=cp.id) AS total_amount,
                  (SELECT MIN(pi.stage) FROM contractor_po_items pi WHERE pi.po_id=cp.id) AS min_stage,
                  (SELECT GROUP_CONCAT(DISTINCT b.name) FROM contractor_po_items pi JOIN order_items oi ON oi.id=pi.order_item_id JOIN orders o ON o.id=oi.order_id JOIN buyers b ON b.id=o.buyer_id WHERE pi.po_id=cp.id) AS buyers,
                  (SELECT GROUP_CONCAT(DISTINCT o.po_number) FROM contractor_po_items pi JOIN order_items oi ON oi.id=pi.order_item_id JOIN orders o ON o.id=oi.order_id WHERE pi.po_id=cp.id) AS buyer_pos,
                  (SELECT filename FROM documents d WHERE d.ref_type='po' AND d.ref_id=cp.id ORDER BY d.id DESC LIMIT 1) AS pdf
                  FROM contractor_pos cp JOIN contractors c ON c.id=cp.contractor_id WHERE {' AND '.join(where)} ORDER BY cp.po_date DESC, cp.id DESC"""
        out = rows(conn.execute(sql, params))
        for o in out:
            o["total_cbm"] = round(o["total_cbm"] or 0, 3)
            o["total_amount"] = round(o["total_amount"] or 0, 2)
            o["stage_name"] = STAGES[o["min_stage"]] if o["min_stage"] is not None else ""
        return out


@app.get("/api/pos/{id}")
def po_get(id: int):
    with get_db() as conn:
        return _po_full(conn, id)


@app.post("/api/pos")
async def po_create(req: Request):
    data = await req.json()
    cid = data.get("contractor_id")
    items = [i for i in data.get("items", []) if f(i.get("qty")) > 0]
    if not cid:
        err("Select a contractor")
    if not items:
        err("Add at least one item with quantity")
    with get_db() as conn:
        c = row(conn.execute("SELECT * FROM contractors WHERE id=?", (cid,)))
        if not c:
            err("Contractor not found")
        # validate available quantity
        for it in items:
            oi = fetch_order_items(conn, "oi.id=?", (it["order_item_id"],))
            if not oi:
                err("Order item not found")
            oi = oi[0]
            avail = oi["remaining_to_assign"] if c["type"] == "wood" else oi["qty"] - oi["iron_assigned_qty"] if c["type"] == "iron" else oi["qty"]
            if f(it["qty"]) > avail + 1e-9 and not data.get("allow_over"):
                err(f"{oi['code']} ({oi['buyer_name']} {oi['po_number']}): only {avail:g} available to assign, you entered {f(it['qty']):g}")
        po_number = clean(data.get("po_number")) or next_po_number(conn)
        if conn.execute("SELECT 1 FROM contractor_pos WHERE po_number=?", (po_number,)).fetchone():
            err(f"PO number {po_number} already exists")
        conn.execute("INSERT INTO contractor_pos(po_number, contractor_id, po_date, delivery_date, notes) VALUES (?,?,?,?,?)",
                     (po_number, cid, data.get("po_date") or today(), data.get("delivery_date") or None, data.get("notes") or ""))
        pid = conn.execute("SELECT last_insert_rowid()").fetchone()[0]
        for it in items:
            conn.execute("INSERT INTO contractor_po_items(po_id, order_item_id, qty, labour_rate, stage, stage_date, remarks) VALUES (?,?,?,?,?,?,?)",
                         (pid, it["order_item_id"], f(it["qty"]), f(it.get("labour_rate")), 0, data.get("po_date") or today(), it.get("remarks") or ""))
            piid = conn.execute("SELECT last_insert_rowid()").fetchone()[0]
            conn.execute("INSERT INTO stage_log(po_item_id, stage, log_date, note) VALUES (?,?,?,?)", (piid, 0, data.get("po_date") or today(), "PO created"))
            if f(it.get("labour_rate")):
                conn.execute("UPDATE order_items SET labour_rate=? WHERE id=? AND (labour_rate IS NULL OR labour_rate=0)", (f(it.get("labour_rate")), it["order_item_id"]))
                conn.execute("UPDATE products SET labour_rate=? WHERE id=(SELECT product_id FROM order_items WHERE id=?) AND (labour_rate IS NULL OR labour_rate=0)", (f(it.get("labour_rate")), it["order_item_id"]))
            if c["type"] == "iron":
                conn.execute("UPDATE order_items SET iron_required=1, iron_ordered_at=COALESCE(iron_ordered_at, ?) WHERE id=?", (data.get("po_date") or today(), it["order_item_id"]))
        fname = _generate_po_doc(conn, pid)
        po = _po_full(conn, pid)
        po["pdf"] = fname
        return po


@app.put("/api/pos/{id}")
async def po_update(id: int, req: Request):
    data = await req.json()
    with get_db() as conn:
        po = row(conn.execute("SELECT * FROM contractor_pos WHERE id=?", (id,)))
        if not po:
            err("PO not found", 404)
        conn.execute("UPDATE contractor_pos SET po_date=?, delivery_date=?, notes=? WHERE id=?",
                     (data.get("po_date", po["po_date"]), data.get("delivery_date", po["delivery_date"]), data.get("notes", po["notes"]), id))
        for it in data.get("items", []):
            if it.get("id"):
                conn.execute("UPDATE contractor_po_items SET qty=?, labour_rate=?, remarks=? WHERE id=? AND po_id=?",
                             (f(it.get("qty")), f(it.get("labour_rate")), it.get("remarks") or "", it["id"], id))
        for iid in data.get("remove_items", []):
            conn.execute("DELETE FROM contractor_po_items WHERE id=? AND po_id=?", (iid, id))
        if data.get("regenerate", True):
            _generate_po_doc(conn, id)
        return _po_full(conn, id)


@app.delete("/api/pos/{id}")
def po_delete(id: int):
    with get_db() as conn:
        conn.execute("DELETE FROM documents WHERE ref_type='po' AND ref_id=?", (id,))
        conn.execute("DELETE FROM contractor_pos WHERE id=?", (id,))
        return {"ok": True}


@app.post("/api/pos/{id}/pdf")
def po_pdf(id: int):
    with get_db() as conn:
        return {"filename": _generate_po_doc(conn, id)}


# ----------------------------------------------------------------------------- production tracking

@app.get("/api/production")
def production(contractor_id: Optional[int] = None, buyer_id: Optional[int] = None, stage: Optional[int] = None, q: str = "", include_done: int = 0):
    with get_db() as conn:
        where, params = ["1=1"], []
        if contractor_id:
            where.append("cp.contractor_id=?"); params.append(contractor_id)
        if buyer_id:
            where.append("o.buyer_id=?"); params.append(buyer_id)
        if stage is not None:
            where.append("pi.stage=?"); params.append(stage)
        if q:
            where.append("(oi.code LIKE ? OR oi.name LIKE ? OR o.po_number LIKE ? OR cp.po_number LIKE ?)"); params += [f"%{q}%"] * 4
        if not include_done:
            where.append("pi.stage < 11")
        sql = f"""SELECT pi.*, cp.po_number AS contractor_po, cp.po_date, cp.delivery_date, c.name AS contractor_name, c.type AS contractor_type,
                  oi.code, oi.name, oi.size, oi.cbm, oi.finish, oi.qty AS order_qty, o.po_number AS buyer_po, o.order_date, b.name AS buyer_name, p.image,
                  (SELECT GROUP_CONCAT(stage || '|' || COALESCE(log_date,'') || '|' || COALESCE(note,''), ';;') FROM stage_log sl WHERE sl.po_item_id=pi.id ORDER BY sl.id) AS history
                  FROM contractor_po_items pi JOIN contractor_pos cp ON cp.id=pi.po_id JOIN contractors c ON c.id=cp.contractor_id
                  JOIN order_items oi ON oi.id=pi.order_item_id JOIN orders o ON o.id=oi.order_id JOIN buyers b ON b.id=o.buyer_id
                  LEFT JOIN products p ON p.id=oi.product_id
                  WHERE {' AND '.join(where)} ORDER BY pi.stage, o.order_date, cp.po_number, pi.id"""
        out = rows(conn.execute(sql, params))
        for it in out:
            it["stage_name"] = STAGES[it["stage"]] if it["stage"] is not None and 0 <= it["stage"] < len(STAGES) else ""
            it["t_cbm"] = round(it["qty"] * (it["cbm"] or 0), 3)
            it["amount"] = round(it["qty"] * (it["labour_rate"] or 0), 2)
        return out


@app.put("/api/production/{id}/stage")
async def set_stage(id: int, req: Request):
    data = await req.json()
    stage = int(data.get("stage", 0))
    if not 0 <= stage < len(STAGES):
        err("Invalid stage")
    with get_db() as conn:
        d = data.get("date") or today()
        conn.execute("UPDATE contractor_po_items SET stage=?, stage_date=?, remarks=COALESCE(?, remarks) WHERE id=?", (stage, d, data.get("remarks"), id))
        conn.execute("INSERT INTO stage_log(po_item_id, stage, log_date, note) VALUES (?,?,?,?)", (id, stage, d, data.get("note") or ""))
        return {"ok": True, "stage": stage, "stage_name": STAGES[stage]}


@app.put("/api/production/bulk-stage")
async def bulk_stage(req: Request):
    data = await req.json()
    stage = int(data.get("stage", 0))
    d = data.get("date") or today()
    with get_db() as conn:
        for id in data.get("ids", []):
            conn.execute("UPDATE contractor_po_items SET stage=?, stage_date=? WHERE id=?", (stage, d, id))
            conn.execute("INSERT INTO stage_log(po_item_id, stage, log_date, note) VALUES (?,?,?,?)", (id, stage, d, data.get("note") or ""))
        return {"ok": True, "count": len(data.get("ids", []))}


# ----------------------------------------------------------------------------- dispatch

@app.get("/api/dispatches")
def dispatches(buyer_id: Optional[int] = None, order_id: Optional[int] = None, date_from: str = "", date_to: str = ""):
    with get_db() as conn:
        where, params = ["1=1"], []
        if buyer_id:
            where.append("o.buyer_id=?"); params.append(buyer_id)
        if order_id:
            where.append("d.order_id=?"); params.append(order_id)
        if date_from:
            where.append("d.dispatch_date >= ?"); params.append(date_from)
        if date_to:
            where.append("d.dispatch_date <= ?"); params.append(date_to)
        sql = f"""SELECT d.*, o.po_number, b.name AS buyer_name,
                  (SELECT SUM(di.qty) FROM dispatch_items di WHERE di.dispatch_id=d.id) AS qty,
                  (SELECT SUM(di.qty*oi.cbm) FROM dispatch_items di JOIN order_items oi ON oi.id=di.order_item_id WHERE di.dispatch_id=d.id) AS cbm,
                  (SELECT SUM(di.qty*oi.rate) FROM dispatch_items di JOIN order_items oi ON oi.id=di.order_item_id WHERE di.dispatch_id=d.id) AS amount,
                  (SELECT GROUP_CONCAT(oi.code || ' x' || CAST(di.qty AS INTEGER), ', ') FROM dispatch_items di JOIN order_items oi ON oi.id=di.order_item_id WHERE di.dispatch_id=d.id) AS items
                  FROM dispatches d JOIN orders o ON o.id=d.order_id JOIN buyers b ON b.id=o.buyer_id
                  WHERE {' AND '.join(where)} ORDER BY d.dispatch_date DESC, d.id DESC"""
        out = rows(conn.execute(sql, params))
        for d in out:
            d["cbm"] = round(d["cbm"] or 0, 3)
            d["amount"] = round(d["amount"] or 0, 2)
        return out


@app.post("/api/dispatches")
async def dispatch_create(req: Request):
    data = await req.json()
    items = [i for i in data.get("items", []) if f(i.get("qty")) > 0]
    if not items:
        err("Enter dispatched quantity for at least one item")
    with get_db() as conn:
        first = fetch_order_items(conn, "oi.id=?", (items[0]["order_item_id"],))
        if not first:
            err("Order item not found")
        order_id = data.get("order_id") or first[0]["order_id"]
        for it in items:
            oi = fetch_order_items(conn, "oi.id=?", (it["order_item_id"],))[0]
            if oi["order_id"] != order_id:
                err("All items of one dispatch must belong to the same buyer PO")
            if f(it["qty"]) > oi["remaining_qty"] + 1e-9:
                err(f"{oi['code']}: only {oi['remaining_qty']:g} remaining, you entered {f(it['qty']):g}")
        conn.execute("INSERT INTO dispatches(order_id, dispatch_date, invoice_no, container, notes) VALUES (?,?,?,?,?)",
                     (order_id, data.get("dispatch_date") or today(), clean(data.get("invoice_no")), clean(data.get("container")), data.get("notes") or ""))
        did = conn.execute("SELECT last_insert_rowid()").fetchone()[0]
        for it in items:
            conn.execute("INSERT INTO dispatch_items(dispatch_id, order_item_id, qty) VALUES (?,?,?)", (did, it["order_item_id"], f(it["qty"])))
            oi = fetch_order_items(conn, "oi.id=?", (it["order_item_id"],))[0]
            if oi["remaining_qty"] <= 0:
                # fully dispatched: move all its production lines to Final Dispatch
                for pi in rows(conn.execute("SELECT id FROM contractor_po_items WHERE order_item_id=? AND stage < 11", (it["order_item_id"],))):
                    conn.execute("UPDATE contractor_po_items SET stage=11, stage_date=? WHERE id=?", (data.get("dispatch_date") or today(), pi["id"]))
                    conn.execute("INSERT INTO stage_log(po_item_id, stage, log_date, note) VALUES (?,?,?,?)", (pi["id"], 11, data.get("dispatch_date") or today(), "Auto: item dispatched"))
        # close order if everything is dispatched
        o = _order_full(conn, order_id)
        conn.execute("UPDATE orders SET status=? WHERE id=?", ("closed" if o["remaining_qty"] <= 0 else "open", order_id))
        return {"ok": True, "dispatch_id": did, "order": o}


@app.delete("/api/dispatches/{id}")
def dispatch_delete(id: int):
    with get_db() as conn:
        d = row(conn.execute("SELECT * FROM dispatches WHERE id=?", (id,)))
        if not d:
            err("Not found", 404)
        conn.execute("DELETE FROM dispatches WHERE id=?", (id,))
        conn.execute("UPDATE orders SET status='open' WHERE id=?", (d["order_id"],))
        return {"ok": True}


# ----------------------------------------------------------------------------- reports

def _report_rows(conn, kind, params):
    """Return (title, subtitle, columns, rows, totals) for the given report kind."""
    cur = get_settings(conn).get("currency", "USD")
    lcur = get_settings(conn).get("labour_currency", "INR")
    where, p = ["oi.qty > 0"], []
    if params.get("buyer_id"):
        where.append("o.buyer_id=?"); p.append(params["buyer_id"])
    if params.get("po"):
        where.append("o.po_number LIKE ?"); p.append(f"%{params['po']}%")
    if params.get("date_from"):
        where.append("o.order_date >= ?"); p.append(params["date_from"])
    if params.get("date_to"):
        where.append("o.order_date <= ?"); p.append(params["date_to"])
    if params.get("code"):
        where.append("oi.code LIKE ?"); p.append(f"%{params['code']}%")
    sub = []
    if params.get("buyer_id"):
        sub.append("Buyer: " + (row(conn.execute("SELECT name FROM buyers WHERE id=?", (params["buyer_id"],))) or {}).get("name", ""))
    if params.get("po"):
        sub.append("PO: " + params["po"])
    if params.get("date_from") or params.get("date_to"):
        sub.append(f"Date: {pdf.fmt_date(params.get('date_from')) or '...'} to {pdf.fmt_date(params.get('date_to')) or '...'}")
    if params.get("code"):
        sub.append("Code: " + params["code"])
    subtitle = " | ".join(sub)
    N, Q = pdf.num, pdf.qty_fmt

    if kind == "orders_by_code":
        items = fetch_order_items(conn, " AND ".join(where), p)
        agg = {}
        for i in items:
            a = agg.setdefault(i["code"], {"code": i["code"], "name": i["name"], "size": i["size"], "cbm": i["cbm"], "rate": i["rate"], "qty": 0, "dispatched": 0, "remaining": 0, "t_cbm": 0, "amount": 0, "buyers": set()})
            a["qty"] += i["qty"]; a["dispatched"] += i["dispatched_qty"]; a["remaining"] += i["remaining_qty"]
            a["t_cbm"] += i["t_cbm"]; a["amount"] += i["amount"]; a["buyers"].add(i["buyer_name"])
        data = []
        tot = {"qty": 0, "dispatched": 0, "remaining": 0, "t_cbm": 0, "amount": 0}
        for a in sorted(agg.values(), key=lambda x: x["code"]):
            for k in tot:
                tot[k] += a[k]
            data.append({**a, "buyers": ", ".join(sorted(a["buyers"])), "qty": Q(a["qty"]), "dispatched": Q(a["dispatched"]), "remaining": Q(a["remaining"]),
                         "cbm": N(a["cbm"], 3), "t_cbm": N(a["t_cbm"], 3), "rate": N(a["rate"]), "amount": N(a["amount"])})
        cols = [("Code", "code", "l", 30), ("Description", "name", "l", 70), ("Size", "size", "l", 34), ("Buyers", "buyers", "l", 34), ("Ordered", "qty", "r", 16),
                ("Dispatched", "dispatched", "r", 18), ("Remaining", "remaining", "r", 18), ("CBM", "cbm", "r", 14), ("T.CBM", "t_cbm", "r", 16), (f"Rate", "rate", "r", 14), (f"Amount {cur}", "amount", "r", 20)]
        totals = {"code": "GRAND TOTAL", "qty": Q(tot["qty"]), "dispatched": Q(tot["dispatched"]), "remaining": Q(tot["remaining"]), "t_cbm": N(tot["t_cbm"], 3), "amount": N(tot["amount"])}
        return "Total Orders - Product Code Wise", subtitle, cols, data, totals

    if kind in ("orders_by_buyer", "remaining_by_buyer"):
        remaining_only = kind == "remaining_by_buyer"
        items = fetch_order_items(conn, " AND ".join(where), p)
        if remaining_only:
            items = [i for i in items if i["remaining_qty"] > 0]
        groups = {}
        for i in items:
            groups.setdefault(i["buyer_name"], {}).setdefault(i["po_number"], []).append(i)
        data = []
        g = {"qty": 0, "cbm": 0, "amount": 0, "labour": 0}
        for buyer in sorted(groups):
            bt = {"qty": 0, "cbm": 0, "amount": 0, "labour": 0}
            data.append({"_group": True, "po": f"BUYER: {buyer}"})
            for po in groups[buyer]:
                pt = {"qty": 0, "cbm": 0, "amount": 0, "labour": 0}
                lst = groups[buyer][po]
                for i in lst:
                    q = i["remaining_qty"] if remaining_only else i["qty"]
                    cb = q * (i["cbm"] or 0); am = q * (i["rate"] or 0); lb = q * (i["labour_rate"] or 0)
                    pt["qty"] += q; pt["cbm"] += cb; pt["amount"] += am; pt["labour"] += lb
                    data.append({"buyer": buyer, "po": po, "date": pdf.fmt_date(i["order_date"]), "code": i["code"], "name": i["name"], "size": i["size"],
                                 "ordered": Q(i["qty"]), "dispatched": Q(i["dispatched_qty"]), "qty": Q(q), "cbm": N(i["cbm"], 3), "t_cbm": N(cb, 3), "rate": N(i["rate"]),
                                 "amount": N(am), "labour_rate": N(i["labour_rate"]), "labour": N(lb), "contractors": i["contractors"] or "", "stage": i["stage_name"], "status": i["item_status"]})
                data.append({"_bold": True, "buyer": "", "po": f"PO {po} total", "qty": Q(pt["qty"]), "t_cbm": N(pt["cbm"], 3), "amount": N(pt["amount"]), "labour": N(pt["labour"])})
                for k in bt:
                    bt[k] += pt[k]
            data.append({"_bold": True, "po": f"{buyer} TOTAL", "qty": Q(bt["qty"]), "t_cbm": N(bt["cbm"], 3), "amount": N(bt["amount"]), "labour": N(bt["labour"])})
            for k in g:
                g[k] += bt[k]
        qlabel = "Remaining" if remaining_only else "Qty"
        cols = [("PO", "po", "l", 30), ("Date", "date", "c", 20), ("Code", "code", "l", 28), ("Description", "name", "l", 52), ("Size", "size", "l", 30),
                ("Ordered", "ordered", "r", 15), ("Disp.", "dispatched", "r", 13), (qlabel, "qty", "r", 16), ("CBM", "cbm", "r", 14), ("T.CBM", "t_cbm", "r", 16),
                ("Rate", "rate", "r", 14), (f"Amount {cur}", "amount", "r", 20), ("Labour", "labour_rate", "r", 14), (f"Labour {lcur}", "labour", "r", 20)]
        if remaining_only:
            cols.append(("Contractors / Stage", "contractors", "l", 30))
            for d in data:
                if "contractors" in d:
                    d["contractors"] = (d["contractors"] + " - " + d["stage"]).strip(" -")
        totals = {"po": "GRAND TOTAL", "qty": Q(g["qty"]), "t_cbm": N(g["cbm"], 3), "amount": N(g["amount"]), "labour": N(g["labour"])}
        title = "Remaining To Dispatch - Buyer & PO Wise" if remaining_only else "Total Orders - Buyer & PO Wise"
        return title, subtitle, cols, data, totals

    if kind == "buyer_summary":
        items = fetch_order_items(conn, " AND ".join(where), p)
        agg = {}
        for i in items:
            a = agg.setdefault(i["buyer_name"], {"buyer": i["buyer_name"], "pos": set(), "lines": 0, "qty": 0, "dispatched": 0, "remaining": 0, "cbm": 0, "rem_cbm": 0, "amount": 0, "rem_amount": 0})
            a["pos"].add(i["po_number"]); a["lines"] += 1; a["qty"] += i["qty"]; a["dispatched"] += i["dispatched_qty"]; a["remaining"] += i["remaining_qty"]
            a["cbm"] += i["t_cbm"]; a["rem_cbm"] += i["remaining_cbm"]; a["amount"] += i["amount"]; a["rem_amount"] += i["remaining_amount"]
        data, tot = [], {"pos": 0, "lines": 0, "qty": 0, "dispatched": 0, "remaining": 0, "cbm": 0, "rem_cbm": 0, "amount": 0, "rem_amount": 0}
        for a in sorted(agg.values(), key=lambda x: x["buyer"]):
            a["pos"] = len(a["pos"])
            for k in tot:
                tot[k] += a[k]
            data.append({"buyer": a["buyer"], "pos": a["pos"], "lines": a["lines"], "qty": Q(a["qty"]), "dispatched": Q(a["dispatched"]), "remaining": Q(a["remaining"]),
                         "cbm": N(a["cbm"], 3), "rem_cbm": N(a["rem_cbm"], 3), "amount": N(a["amount"]), "rem_amount": N(a["rem_amount"])})
        cols = [("Buyer", "buyer", "l", 40), ("POs", "pos", "r", 14), ("Lines", "lines", "r", 14), ("Ordered", "qty", "r", 18), ("Dispatched", "dispatched", "r", 18), ("Remaining", "remaining", "r", 18),
                ("Total CBM", "cbm", "r", 20), ("Rem. CBM", "rem_cbm", "r", 20), (f"Amount {cur}", "amount", "r", 24), (f"Rem. Amount {cur}", "rem_amount", "r", 26)]
        totals = {"buyer": "GRAND TOTAL", "pos": tot["pos"], "lines": tot["lines"], "qty": Q(tot["qty"]), "dispatched": Q(tot["dispatched"]), "remaining": Q(tot["remaining"]),
                  "cbm": N(tot["cbm"], 3), "rem_cbm": N(tot["rem_cbm"], 3), "amount": N(tot["amount"]), "rem_amount": N(tot["rem_amount"])}
        return "Buyer Summary", subtitle, cols, data, totals

    if kind == "contractor":
        w, pp = ["1=1"], []
        if params.get("contractor_id"):
            w.append("cp.contractor_id=?"); pp.append(params["contractor_id"])
            sub.insert(0, "Contractor: " + (row(conn.execute("SELECT name FROM contractors WHERE id=?", (params["contractor_id"],))) or {}).get("name", ""))
        if params.get("buyer_id"):
            w.append("o.buyer_id=?"); pp.append(params["buyer_id"])
        if params.get("po"):
            w.append("(o.po_number LIKE ? OR cp.po_number LIKE ?)"); pp += [f"%{params['po']}%"] * 2
        if params.get("date_from"):
            w.append("cp.po_date >= ?"); pp.append(params["date_from"])
        if params.get("date_to"):
            w.append("cp.po_date <= ?"); pp.append(params["date_to"])
        if params.get("code"):
            w.append("oi.code LIKE ?"); pp.append(f"%{params['code']}%")
        if params.get("pending_only"):
            w.append("pi.stage < 11")
        lines = rows(conn.execute(f"""SELECT pi.*, cp.po_number AS cpo, cp.po_date, cp.delivery_date, c.name AS contractor_name, oi.code, oi.name, oi.size, oi.cbm,
                                      o.po_number AS buyer_po, b.name AS buyer_name FROM contractor_po_items pi JOIN contractor_pos cp ON cp.id=pi.po_id
                                      JOIN contractors c ON c.id=cp.contractor_id JOIN order_items oi ON oi.id=pi.order_item_id JOIN orders o ON o.id=oi.order_id
                                      JOIN buyers b ON b.id=o.buyer_id WHERE {' AND '.join(w)} ORDER BY c.name, cp.po_date, cp.po_number, pi.id""", pp))
        data, g = [], {"qty": 0, "cbm": 0, "amount": 0}
        groups = {}
        for l in lines:
            groups.setdefault(l["contractor_name"], []).append(l)
        for cname in sorted(groups):
            data.append({"_group": True, "cpo": f"CONTRACTOR: {cname}"})
            ct = {"qty": 0, "cbm": 0, "amount": 0}
            for l in groups[cname]:
                cb = l["qty"] * (l["cbm"] or 0); am = l["qty"] * (l["labour_rate"] or 0)
                ct["qty"] += l["qty"]; ct["cbm"] += cb; ct["amount"] += am
                data.append({"contractor": cname, "cpo": l["cpo"], "date": pdf.fmt_date(l["po_date"]), "buyer": l["buyer_name"], "buyer_po": l["buyer_po"], "code": l["code"], "name": l["name"],
                             "size": l["size"], "qty": Q(l["qty"]), "cbm": N(l["cbm"], 3), "t_cbm": N(cb, 3), "rate": N(l["labour_rate"]), "amount": N(am),
                             "stage": STAGES[l["stage"]] if l["stage"] is not None else "", "delivery": pdf.fmt_date(l["delivery_date"]), "remarks": l["remarks"] or ""})
            data.append({"_bold": True, "cpo": f"{cname} TOTAL", "qty": Q(ct["qty"]), "t_cbm": N(ct["cbm"], 3), "amount": N(ct["amount"])})
            for k in g:
                g[k] += ct[k]
        cols = [("Our PO", "cpo", "l", 24), ("Date", "date", "c", 20), ("Buyer", "buyer", "l", 24), ("Buyer PO", "buyer_po", "l", 24),
                ("Code", "code", "l", 26), ("Description", "name", "l", 46), ("Size", "size", "l", 28), ("Qty", "qty", "r", 13), ("CBM", "cbm", "r", 14), ("T.CBM", "t_cbm", "r", 16),
                (f"Labour Rate", "rate", "r", 17), (f"Amount {lcur}", "amount", "r", 22), ("Stage", "stage", "l", 22), ("Delivery", "delivery", "c", 20)]
        totals = {"cpo": "GRAND TOTAL", "qty": Q(g["qty"]), "t_cbm": N(g["cbm"], 3), "amount": N(g["amount"])}
        return "Contractor Orders", " | ".join(sub), cols, data, totals

    if kind == "production":
        lines = production(contractor_id=params.get("contractor_id"), buyer_id=params.get("buyer_id"), q=params.get("code") or "")
        data = [{"stage": l["stage_name"], "contractor": l["contractor_name"], "cpo": l["contractor_po"], "buyer": l["buyer_name"], "buyer_po": l["buyer_po"], "code": l["code"], "name": l["name"],
                 "qty": Q(l["qty"]), "t_cbm": N(l["t_cbm"], 3), "date": pdf.fmt_date(l["stage_date"]), "delivery": pdf.fmt_date(l["delivery_date"]), "remarks": l["remarks"] or ""} for l in lines]
        cols = [("Stage", "stage", "l", 26), ("Contractor", "contractor", "l", 24), ("Our PO", "cpo", "l", 20), ("Buyer", "buyer", "l", 26), ("Buyer PO", "buyer_po", "l", 22), ("Code", "code", "l", 26),
                ("Description", "name", "l", 56), ("Qty", "qty", "r", 12), ("T.CBM", "t_cbm", "r", 14), ("Stage date", "date", "c", 18), ("Delivery", "delivery", "c", 18), ("Remarks", "remarks", "l", 30)]
        return "Production Status", " | ".join(sub), cols, data, None

    err("Unknown report")


@app.get("/api/reports/{kind}")
def report(kind: str, request: Request, format: str = "json"):
    params = dict(request.query_params)
    for k in ("buyer_id", "contractor_id"):
        if params.get(k):
            params[k] = int(params[k])
    with get_db() as conn:
        title, subtitle, cols, data, totals = _report_rows(conn, kind, params)
        if format == "json":
            return {"title": title, "subtitle": subtitle, "columns": [{"label": c[0], "key": c[1], "align": c[2]} for c in cols], "rows": data, "totals": totals}
        if format == "pdf":
            fname = pdf.report_pdf(title, subtitle, cols, data, get_settings(conn), totals=totals)
            return {"filename": fname}
        if format == "xlsx":
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = title[:30]
            ws.append([title]); ws["A1"].font = Font(bold=True, size=14)
            ws.append([subtitle]); ws.append([])
            ws.append([c[0] for c in cols])
            for cell in ws[4]:
                cell.font = Font(bold=True); cell.fill = PatternFill("solid", fgColor="DDDDDD")
            for r in data:
                ws.append([r.get(c[1], "") for c in cols])
                if r.get("_group") or r.get("_bold"):
                    for cell in ws[ws.max_row]:
                        cell.font = Font(bold=True)
                        cell.fill = PatternFill("solid", fgColor="DFE8F5" if r.get("_group") else "F3F3F3")
            if totals:
                ws.append([totals.get(c[1], "") for c in cols])
                for cell in ws[ws.max_row]:
                    cell.font = Font(bold=True); cell.fill = PatternFill("solid", fgColor="FFE9B3")
            for i, c in enumerate(cols, 1):
                ws.column_dimensions[get_column_letter(i)].width = max(10, c[3] * 0.55)
                for cell in ws[get_column_letter(i)]:
                    if c[2] == "r":
                        cell.alignment = Alignment(horizontal="right")
            buf = io.BytesIO(); wb.save(buf); buf.seek(0)
            fname = pdf.safe_name(title) + " " + datetime.now().strftime("%Y-%m-%d") + ".xlsx"
            return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="{fname}"'})
    err("Unknown format")


@app.get("/api/documents")
def documents(limit: int = 100):
    with get_db() as conn:
        return rows(conn.execute("SELECT * FROM documents ORDER BY id DESC LIMIT ?", (limit,)))


@app.get("/open/{filename}")
def open_pdf(filename: str):
    path = os.path.join(PDF_DIR, filename)
    if not os.path.exists(path):
        err("File not found", 404)
    return FileResponse(path, media_type="application/pdf", filename=filename)


# ----------------------------------------------------------------------------- backup

def make_backup(keep=60):
    """Copy the database into data/backups and prune old copies."""
    bdir = os.path.join(os.path.dirname(PDF_DIR), "backups")
    os.makedirs(bdir, exist_ok=True)
    name = f"oms-backup-{datetime.now().strftime('%Y%m%d-%H%M%S')}.db"
    with get_db() as conn:
        dest = __import__("sqlite3").connect(os.path.join(bdir, name))
        conn.backup(dest)
        dest.close()
    old = sorted(f for f in os.listdir(bdir) if f.startswith("oms-backup-"))
    for f in old[:-keep]:
        try:
            os.remove(os.path.join(bdir, f))
        except OSError:
            pass
    return os.path.join(bdir, name)


@app.post("/api/backup")
def backup():
    return {"ok": True, "file": make_backup()}


def _auto_backup_loop():
    import time
    while True:
        time.sleep(24 * 3600)
        try:
            make_backup()
        except Exception as e:  # never let the backup thread die silently
            print("auto backup failed:", e)


threading.Thread(target=_auto_backup_loop, daemon=True).start()


if __name__ == "__main__":
    import uvicorn
    port = int(os.environ.get("OMS_PORT", "8765"))
    host = os.environ.get("OMS_HOST", "0.0.0.0")
    if os.environ.get("OMS_NO_BROWSER") != "1":
        threading.Timer(1.2, lambda: webbrowser.open(f"http://127.0.0.1:{port}")).start()
    try:
        import socket
        lan_ip = socket.gethostbyname(socket.gethostname())
        print(f"Order Management System running.  This PC: http://127.0.0.1:{port}   Other devices on your network: http://{lan_ip}:{port}")
    except Exception:
        pass
    uvicorn.run(app, host=host, port=port, log_level="warning", proxy_headers=True, forwarded_allow_ips="*")
